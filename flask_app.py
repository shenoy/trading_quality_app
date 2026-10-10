import csv
import io
import itertools
import json
import math
import re
from datetime import date as date_type, datetime, time as time_type, timedelta
from flask import Flask, jsonify, redirect, render_template, request, send_file, url_for
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import text

app = Flask(__name__)

# Configure SQLite database (PythonAnywhere path compatible)
app.config['SQLALCHEMY_DATABASE_URI'] = 'sqlite:///database.db'
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
db = SQLAlchemy(app)


# --- DATABASE MODELS ---


class Trade(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    position = db.Column(db.Integer, default=0)
    setup_quality_str = db.Column(db.String(20), nullable=False)
    ticker = db.Column(db.String(20), nullable=False)
    entry_datetime = db.Column(db.String(50), nullable=False)
    profit_taken_str = db.Column(db.String(10), nullable=False)
    stop_loss_zero_str = db.Column(db.String(10), nullable=False)
    current_equity = db.Column(db.Float, nullable=False)
    risk_str = db.Column(db.String(20), nullable=False)
    risk_reward_str = db.Column(db.String(20), nullable=False)

    risk_under_half_str = db.Column(db.String(10), nullable=True)
    daily_risk_str = db.Column(db.String(10), nullable=True)

    overtrading_str = db.Column(db.String(10), nullable=True, default='no')
    revenge_trading_str = db.Column(db.String(10), nullable=True, default='no')
    impatient_trading_str = db.Column(db.String(10), nullable=True, default='no')
    fearful_trading_str = db.Column(db.String(10), nullable=True, default='no')

    setup_score = db.Column(db.Float, nullable=False)
    profit_score = db.Column(db.Float, nullable=False)
    stop_loss_score = db.Column(db.Float, nullable=False)
    risk_score = db.Column(db.Float, nullable=False)
    risk_reward_score = db.Column(db.Float, nullable=False)
    daily_risk_score = db.Column(db.Float, nullable=True, default=0.0)
    psychology_score = db.Column(db.Float, nullable=True, default=0.0)
    total_score = db.Column(db.Float, nullable=False)
    timestamp = db.Column(db.DateTime, default=datetime.utcnow)
    screenshot_url = db.Column(db.String(500), nullable=True)
    account_id = db.Column(db.Integer, db.ForeignKey('account.id'), nullable=True)
    # When the trade was closed, e.g. 2026-10-09T13:30:05. Optional: the
    # equity chart is drawn in exit-time order and falls back to the entry time.
    exit_datetime = db.Column(db.String(50), nullable=True)

    @property
    def psychology_flags(self):
        flags = []
        if self.overtrading_str == 'yes':
            flags.append('Overtrading')
        if self.revenge_trading_str == 'yes':
            flags.append('Revenge')
        if self.impatient_trading_str == 'yes':
            flags.append('Impatient')
        if self.fearful_trading_str == 'yes':
            flags.append('Fearful')
        return flags


class Account(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(60), unique=True, nullable=False)
    # The balance this account started from. Optional (blank = not set).
    starting_balance = db.Column(db.Float, nullable=True)
    created = db.Column(db.DateTime, default=datetime.utcnow)


class IBKRTrade(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    ticker = db.Column(db.String(20), nullable=False)
    time_entered = db.Column(db.String(50), nullable=False)
    time_exited = db.Column(db.String(50), nullable=False)
    entry_price = db.Column(db.Float, nullable=False, default=0.0)
    exit_price = db.Column(db.Float, nullable=False, default=0.0)
    pnl = db.Column(db.Float, nullable=False, default=0.0)
    balance = db.Column(db.Float, nullable=False, default=0.0)
    timestamp = db.Column(db.DateTime, default=datetime.utcnow)


class AppSetting(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    key = db.Column(db.String(50), unique=True, nullable=False)
    value = db.Column(db.String(200), nullable=False)


with app.app_context():
    db.create_all()
    if not AppSetting.query.filter_by(key='starting_equity').first():
        db.session.add(AppSetting(key='starting_equity', value='10000'))
    if not AppSetting.query.filter_by(key='default_risk_limit').first():
        db.session.add(AppSetting(key='default_risk_limit', value='20'))
    db.session.commit()


# --- SCORING HELPER LOGIC ---

SCORING_RULES = {
    'setup_good': {
        'count': 3,
        'defaults': [1.0, 0.5, 0.0],
        'hint': 'Points for: Good / Medium / Bad',
    },
    'profit_yes': {
        'count': 2,
        'defaults': [1.0, 0.0],
        'hint': 'Points for: Yes / No',
    },
    'sl_yes': {'count': 2, 'defaults': [1.0, 0.0], 'hint': 'Points for: Yes / No'},
    'risk_under_half': {
        'count': 2,
        'defaults': [5.0, -5.0],
        'hint': 'Points for "Risk under 0.5%?": Yes / No',
    },
    'daily_dd': {
        'count': 2,
        'defaults': [5.0, -5.0],
        'hint': 'Points for "Daily drawdown under 1.5%?": Yes / No',
    },
    'rr_bands': {
        'count': 6,
        'defaults': [-5.0, 1.0, 5.0, 15.0, 100.0, 200.0],
        'hint': (
            'Points for R:R in this order: below 0 / 0 to 0.5 / 0.5 to 1 / '
            '1 to 2 / 2 to 3 / 3 and over'
        ),
    },
    'overtrading_yes': {
        'count': 2,
        'defaults': [-1.0, 0.0],
        'hint': 'Points for "Overtrading?": Yes / No',
    },
    'revenge_yes': {
        'count': 2,
        'defaults': [-1.0, 0.0],
        'hint': 'Points for "Revenge trading?": Yes / No',
    },
    'impatient_yes': {
        'count': 2,
        'defaults': [-1.0, 0.0],
        'hint': 'Points for "Impatient trading?": Yes / No',
    },
    'fearful_yes': {
        'count': 2,
        'defaults': [-1.0, 0.0],
        'hint': 'Points for "Fearful trading?": Yes / No',
    },
}


def parse_points(text, count):
    parts = [p for p in re.split(r'[/,\s]+', str(text or '').strip()) if p]
    if len(parts) != count:
        return None
    try:
        return [float(p) for p in parts]
    except ValueError:
        return None


def format_points(points):
    return ' / '.join(f'{p:g}' for p in points)


def get_rule_points(key):
    rule = SCORING_RULES[key]
    setting = AppSetting.query.filter_by(key=key).first()
    if setting:
        parsed = parse_points(setting.value, rule['count'])
        if parsed is not None:
            return parsed
    return list(rule['defaults'])


def to_float_or_none(value):
    try:
        return float(value)
    except (ValueError, TypeError):
        return None


def yes_no_points(answer, pts):
    answer = (answer or '').lower()
    if answer == 'yes':
        return pts[0]
    if answer == 'no':
        return pts[1]
    return 0.0


def calculate_scores(data):
    setup_pts = get_rule_points('setup_good')
    profit_pts = get_rule_points('profit_yes')
    sl_pts = get_rule_points('sl_yes')
    risk_pts = get_rule_points('risk_under_half')
    daily_pts = get_rule_points('daily_dd')
    rr_pts = get_rule_points('rr_bands')

    sq = (data.get('setup_quality') or '').lower()
    if sq == 'good':
        setup_score = setup_pts[0]
    elif sq == 'medium':
        setup_score = setup_pts[1]
    elif sq == 'bad':
        setup_score = setup_pts[2]
    else:
        setup_score = 0.0

    pt = (data.get('profit_taken') or '').lower()
    profit_score = profit_pts[0] if pt == 'yes' else profit_pts[1]

    sl = (data.get('stop_loss_zero') or '').lower()
    stop_loss_score = sl_pts[0] if sl == 'yes' else sl_pts[1]

    risk_score = yes_no_points(data.get('risk_under_half'), risk_pts)
    daily_risk_score = yes_no_points(data.get('daily_risk'), daily_pts)

    rr_val = to_float_or_none(data.get('risk_reward_val'))
    if rr_val is None:
        rr_score = 0.0
    elif rr_val < 0:
        rr_score = rr_pts[0]
    elif rr_val < 0.5:
        rr_score = rr_pts[1]
    elif rr_val < 1:
        rr_score = rr_pts[2]
    elif rr_val < 2:
        rr_score = rr_pts[3]
    elif rr_val < 3:
        rr_score = rr_pts[4]
    else:
        rr_score = rr_pts[5]

    def psych(field, rule_key):
        return yes_no_points(data.get(field) or 'no', get_rule_points(rule_key))

    psychology_score = (
        psych('overtrading', 'overtrading_yes')
        + psych('revenge_trading', 'revenge_yes')
        + psych('impatient_trading', 'impatient_yes')
        + psych('fearful_trading', 'fearful_yes')
    )

    total_score = (
        setup_score
        + profit_score
        + stop_loss_score
        + risk_score
        + daily_risk_score
        + rr_score
        + psychology_score
    )

    return {
        'setup_score': setup_score,
        'profit_score': profit_score,
        'stop_loss_score': stop_loss_score,
        'risk_score': risk_score,
        'daily_risk_score': daily_risk_score,
        'psychology_score': psychology_score,
        'risk_reward_score': rr_score,
        'total_score': round(total_score, 4),
    }


def apply_scores(trade, scores):
    trade.setup_score = scores['setup_score']
    trade.profit_score = scores['profit_score']
    trade.stop_loss_score = scores['stop_loss_score']
    trade.risk_score = scores['risk_score']
    trade.daily_risk_score = scores['daily_risk_score']
    trade.psychology_score = scores['psychology_score']
    trade.risk_reward_score = scores['risk_reward_score']
    trade.total_score = scores['total_score']


def recalculate_all_trades():
    for trade in Trade.query.all():
        apply_scores(
            trade,
            calculate_scores({
                'setup_quality': trade.setup_quality_str,
                'profit_taken': trade.profit_taken_str,
                'stop_loss_zero': trade.stop_loss_zero_str,
                'risk_under_half': trade.risk_under_half_str,
                'daily_risk': trade.daily_risk_str,
                'risk_reward_val': trade.risk_reward_str,
                'overtrading': trade.overtrading_str,
                'revenge_trading': trade.revenge_trading_str,
                'impatient_trading': trade.impatient_trading_str,
                'fearful_trading': trade.fearful_trading_str,
            }),
        )
    db.session.commit()


def ensure_schema():
    columns = {
        row[1] for row in db.session.execute(text('PRAGMA table_info(trade)'))
    }
    changed = False
    try:
        # Every trade belongs to an account. Adding this column on its own
        # doesn't need a re-score, so it doesn't set `changed`.
        if 'account_id' not in columns:
            db.session.execute(
                text('ALTER TABLE trade ADD COLUMN account_id INTEGER')
            )
        # Exit time: also doesn't change any score, so no re-score either.
        if 'exit_datetime' not in columns:
            db.session.execute(
                text('ALTER TABLE trade ADD COLUMN exit_datetime VARCHAR(50)')
            )
        for name, ddl in [
            ('daily_risk_score', 'REAL DEFAULT 0.0'),
            ('psychology_score', 'REAL DEFAULT 0.0'),
            ('overtrading_str', "TEXT DEFAULT 'no'"),
            ('revenge_trading_str', "TEXT DEFAULT 'no'"),
            ('impatient_trading_str', "TEXT DEFAULT 'no'"),
            ('fearful_trading_str', "TEXT DEFAULT 'no'"),
            ('screenshot_url', 'VARCHAR(500)'),
        ]:
            if name not in columns:
                db.session.execute(
                    text(f'ALTER TABLE trade ADD COLUMN {name} {ddl}')
                )
                changed = True
        if 'risk_under_half_str' not in columns:
            db.session.execute(
                text('ALTER TABLE trade ADD COLUMN risk_under_half_str TEXT')
            )
            if 'risk_over_half_str' in columns:
                db.session.execute(
                    text(
                        "UPDATE trade SET risk_under_half_str = CASE"
                        " risk_over_half_str WHEN 'yes' THEN 'no' WHEN 'no'"
                        " THEN 'yes' END"
                    )
                )
            changed = True
        db.session.commit()
    except Exception:
        db.session.rollback()
        app.logger.exception('Schema upgrade skipped')
        return
    if changed:
        recalculate_all_trades()


with app.app_context():
    ensure_schema()


# --- ACCOUNTS ---
# Each trade belongs to an account (your main account, a prop firm challenge,
# ...). The account you are looking at is remembered in a cookie, and every
# trade page and chart only shows that account's trades.


def ensure_default_account():
    """Makes sure one account exists and every trade belongs to an account."""
    try:
        first = Account.query.order_by(Account.id.asc()).first()
        if first is None:
            first = Account(name='Main Account')
            db.session.add(first)
            db.session.flush()
        Trade.query.filter(Trade.account_id.is_(None)).update(
            {'account_id': first.id}, synchronize_session=False
        )
        db.session.commit()
    except Exception:
        db.session.rollback()
        app.logger.exception('Default account setup skipped')


with app.app_context():
    ensure_default_account()


def get_active_account():
    try:
        account = db.session.get(Account, int(request.cookies.get('account_id', '')))
    except ValueError:
        account = None
    return account or Account.query.order_by(Account.id.asc()).first()


def parse_balance(value):
    """Returns (number, error). A blank value means 'not set' -> (None, None)."""
    value = (value or '').strip().replace(',', '')
    if not value:
        return None, None
    try:
        number = float(value)
    except ValueError:
        return None, 'Starting balance must be a number.'
    if number != number or number in (float('inf'), float('-inf')):
        return None, 'Starting balance must be a number.'
    return number, None


def format_number(number):
    return ('%.2f' % number).rstrip('0').rstrip('.')


def validate_account_name(name, exclude_id=None):
    if not name:
        return 'Please enter a name for the account.'
    if len(name) > 60:
        return 'The name can be at most 60 characters.'
    query = Account.query.filter(db.func.lower(Account.name) == name.lower())
    if exclude_id is not None:
        query = query.filter(Account.id != exclude_id)
    if query.first():
        return 'An account with that name already exists.'
    return None


def trade_profit_loss(trade):
    try:
        return float(trade.risk_str) * float(trade.risk_reward_str)
    except (TypeError, ValueError):
        return 0.0


def account_summary(account):
    """Balance progress and trade stats for one account."""
    trades = (
        Trade.query.filter_by(account_id=account.id)
        .order_by(Trade.position.asc())
        .all()
    )
    start = account.starting_balance
    # With no trades yet the balance is still the starting balance.
    latest = trades[-1].current_equity if trades else start
    change = latest - start if latest is not None and start is not None else None
    change_pct = change / start * 100 if change is not None and start and start > 0 else None
    return {
        'trade_count': len(trades),
        'starting_balance': start,
        'latest_equity': latest,
        'change': change,
        'change_pct': change_pct,
        'net_pl': sum(trade_profit_loss(t) for t in trades),
        'avg_score': (
            sum(t.total_score for t in trades) / len(trades) if trades else None
        ),
    }


@app.context_processor
def inject_accounts():
    return {
        'accounts': Account.query.order_by(Account.id.asc()).all(),
        'active_account': get_active_account(),
        'account_summary': account_summary,
    }


def safe_next(target):
    """Only allow redirects to pages inside this app."""
    if (
        target
        and target.startswith('/')
        and not target.startswith('//')
        and '\\' not in target
    ):
        return target
    return '/'


def switch_to(account, target='/'):
    response = redirect(safe_next(target))
    response.set_cookie(
        'account_id', str(account.id), max_age=60 * 60 * 24 * 365, samesite='Lax'
    )
    return response


ACCOUNT_MESSAGES = {
    'deleted': ('success', 'Account deleted.'),
    'has_trades': (
        'warning',
        "That account still has trades, so it can't be deleted. "
        'Delete its trades first.',
    ),
    'last_account': ('warning', "You can't delete your only account."),
}


@app.route('/accounts', methods=['GET', 'POST'])
def accounts_page():
    error = None
    form = {'name': '', 'starting_balance': ''}
    if request.method == 'POST':
        form['name'] = (request.form.get('name') or '').strip()
        form['starting_balance'] = (request.form.get('starting_balance') or '').strip()
        balance, error = parse_balance(form['starting_balance'])
        if error is None and balance is None:
            error = 'Please enter the starting balance for the new account.'
        if error is None:
            error = validate_account_name(form['name'])
        if error is None:
            account = Account(name=form['name'], starting_balance=balance)
            db.session.add(account)
            db.session.commit()
            return switch_to(account, '/')  # open the new account's dashboard
    return render_template(
        'accounts.html',
        error=error,
        form=form,
        message=ACCOUNT_MESSAGES.get(request.args.get('msg')),
    )


@app.route('/accounts/edit/<int:id>', methods=['GET', 'POST'])
def edit_account(id):
    account = Account.query.get_or_404(id)
    error = None
    form = {
        'name': account.name,
        'starting_balance': (
            ''
            if account.starting_balance is None
            else format_number(account.starting_balance)
        ),
    }
    if request.method == 'POST':
        form['name'] = (request.form.get('name') or '').strip()
        form['starting_balance'] = (request.form.get('starting_balance') or '').strip()
        balance, error = parse_balance(form['starting_balance'])
        if error is None:
            error = validate_account_name(form['name'], exclude_id=account.id)
        if error is None:
            account.name = form['name']
            account.starting_balance = balance
            db.session.commit()
            return redirect(url_for('accounts_page'))
    return render_template('edit_account.html', account=account, error=error, form=form)


@app.route('/accounts/delete/<int:id>', methods=['POST'])
def delete_account(id):
    account = Account.query.get_or_404(id)
    if Account.query.count() <= 1:
        return redirect(url_for('accounts_page', msg='last_account'))
    if Trade.query.filter_by(account_id=account.id).first():
        return redirect(url_for('accounts_page', msg='has_trades'))
    db.session.delete(account)
    db.session.commit()
    return redirect(url_for('accounts_page', msg='deleted'))


@app.route('/accounts/switch/<int:id>')
def switch_account(id):
    account = Account.query.get_or_404(id)
    return switch_to(account, request.args.get('next', '/'))


# --- IMPORT FROM EXCEL ---
# Upload an Excel sheet (.xlsx) or a CSV. Two things can come out of it, and
# both only happen after you confirm a preview:
#   1. Trades that have no screenshot yet get their link from the sheet
#      (matched by ticker + entry date + entry time; a saved link is never changed).
#   2. Rows that are in the sheet but not in the app can be added as new trades.
#      The sheet can't answer the app's yes/no questions, so those get defaults:
#      setup quality is guessed from the notes, everything else is "No".

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_SHEET_ROWS = 20000
MAX_SHEET_COLUMNS = 60
MAX_LISTED_ROWS = 300
MAX_NEW_TRADES = 2000

# The preview form sends a few fields per row, so allow big sheets through.
app.config['MAX_FORM_PARTS'] = 20000
app.config['MAX_FORM_MEMORY_SIZE'] = 10 * 1024 * 1024
app.request_class.max_form_parts = 20000
app.request_class.max_form_memory_size = 10 * 1024 * 1024

COLUMN_NAMES = {
    'ticker': 'a ticker column (e.g. GBPCAD)',
    'date': 'a date column',
    'time': 'an entry time column',
    'link': 'a screenshot link column',
}

class UploadError(Exception):
    """A problem with the uploaded file, worded for the person uploading it."""


def column_letter(index):
    letters = ''
    index += 1
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def normalize_ticker(value):
    return re.sub(r'[^A-Z0-9]', '', str(value or '').upper())


TICKER_PATTERN = re.compile(r'^(?=(?:.*[A-Z]){2})[A-Z][A-Z0-9]{2,9}$')


def looks_like_ticker(value):
    return isinstance(value, str) and bool(TICKER_PATTERN.match(normalize_ticker(value)))


def clean_url(value, hyperlink=None):
    """A usable http(s) link from a cell's text or its hyperlink, else None."""
    for candidate in (value, hyperlink):
        link = str(candidate).strip() if candidate is not None else ''
        if (
            link.lower().startswith(('http://', 'https://'))
            and not re.search(r'\s', link)
            and len(link) <= 500
        ):
            return link
    return None


DATE_FORMATS = [
    '%d-%b-%Y', '%d-%b', '%d %b %Y', '%d %b', '%d-%B-%Y', '%d-%B',
    '%d %B %Y', '%d %B', '%d/%m/%Y', '%d/%m/%y', '%Y-%m-%d', '%d.%m.%Y',
]
TIME_PATTERN = re.compile(r'^(\d{1,2}):(\d{2})(?::\d{2})?\s*([AaPp][Mm])?$')


def parse_date_cell(value):
    """(year or None, month, day) from a date cell, or None."""
    if isinstance(value, datetime):
        if value.year < 1990:  # a time-only value stored as a datetime
            return None
        return value.year, value.month, value.day
    if isinstance(value, date_type):
        return value.year, value.month, value.day
    if isinstance(value, str):
        text_value = value.strip()
        for fmt in DATE_FORMATS:
            has_year = '%Y' in fmt or '%y' in fmt
            try:
                if has_year:
                    parsed = datetime.strptime(text_value, fmt)
                else:  # "2-Oct" has no year; 2000 is just a placeholder
                    parsed = datetime.strptime(text_value + ' 2000', fmt + ' %Y')
            except ValueError:
                continue
            return (parsed.year if has_year else None), parsed.month, parsed.day
    return None


def parse_time_cell(value):
    """(hour, minute) from a time cell, or None."""
    if isinstance(value, datetime):
        if value.year >= 1990 and value.hour == 0 and value.minute == 0:
            return None  # a plain date
        return value.hour, value.minute
    if isinstance(value, time_type):
        return value.hour, value.minute
    if isinstance(value, timedelta):
        minutes = int(value.total_seconds() // 60)
        return divmod(minutes, 60) if 0 <= minutes < 1440 else None
    if isinstance(value, str):
        match = TIME_PATTERN.match(value.strip())
        if match:
            hour, minute, meridiem = int(match.group(1)), int(match.group(2)), match.group(3)
            if meridiem:
                hour = hour % 12 + (12 if meridiem.lower() == 'pm' else 0)
            if hour < 24 and minute < 60:
                return hour, minute
    return None


def parse_entry_datetime(value):
    """The app stores entry times as text such as 2026-10-06T11:16."""
    value = str(value or '').strip().replace('T', ' ')
    for candidate in (value, value[:16]):
        for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M'):
            try:
                return datetime.strptime(candidate, fmt)
            except ValueError:
                continue
    return None


def format_exit_datetime(when):
    """2026-10-09T13:30:05 (seconds only when there are some)."""
    if when.second:
        return f'{when:%Y-%m-%dT%H:%M:%S}'
    return f'{when:%Y-%m-%dT%H:%M}'


def clean_exit_datetime(value):
    """An exit time in the app's own text format, or None if blank/invalid."""
    when = parse_entry_datetime(value)
    return format_exit_datetime(when) if when else None


EXIT_CELL_FORMATS = (
    '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M',
    '%d/%m/%Y %H:%M:%S.%f', '%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M',
)


def parse_exit_cell(value):
    """A full date + time from an 'Exit time' cell, or None. A bare time or a
    bare date can't say when the trade closed, so those are ignored."""
    if isinstance(value, datetime):
        if value.year < 1990:
            return None  # a time-only value stored as a datetime
        if value.hour == 0 and value.minute == 0 and value.second == 0:
            return None  # a plain date
        return value.replace(microsecond=0)
    if isinstance(value, str):
        text_value = value.strip().replace('T', ' ')
        for fmt in EXIT_CELL_FORMATS:
            try:
                return datetime.strptime(text_value, fmt).replace(microsecond=0)
            except ValueError:
                continue
    return None


def has_screenshot(trade):
    return bool((trade.screenshot_url or '').strip())


def read_sheets(filename, data):
    """Returns [(sheet name, rows)]; each row is a list of (value, hyperlink)."""
    extension = filename.lower().rsplit('.', 1)[-1] if '.' in filename else ''
    if extension in ('xlsx', 'xlsm'):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise UploadError(
                "Reading Excel files needs the 'openpyxl' package. "
                'Install it with:  pip install openpyxl'
            )
        try:
            workbook = load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            raise UploadError(
                "That file couldn't be read as an Excel workbook. "
                'Is it a real .xlsx file?'
            )
        sheets = []
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(
                max_row=min(sheet.max_row or 1, MAX_SHEET_ROWS),
                max_col=min(sheet.max_column or 1, MAX_SHEET_COLUMNS),
            ):
                cells = []
                for cell in row:
                    link = getattr(cell, 'hyperlink', None)
                    cells.append((cell.value, link.target if link else None))
                rows.append(cells)
            sheets.append((sheet.title, rows))
        return sheets
    if extension == 'csv':
        text_data = data.decode('utf-8-sig', errors='replace')
        reader = csv.reader(io.StringIO(text_data))
        rows = [[(cell, None) for cell in row[:MAX_SHEET_COLUMNS]] for row in reader]
        return [('CSV file', rows[:MAX_SHEET_ROWS])]
    if extension == 'xls':
        raise UploadError(
            'That is an old-style .xls file. In Excel use File > Save As > '
            '"Excel Workbook (.xlsx)" and upload that file instead.'
        )
    raise UploadError('Please upload an Excel file (.xlsx) or a .csv file.')


def parse_number(value):
    """A finite number from a cell (also accepts '1,234.5' or '$50'), else None."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    elif isinstance(value, str):
        cleaned = value.strip().replace(',', '').replace('$', '')
        if not cleaned:
            return None
        try:
            number = float(cleaned)
        except ValueError:
            return None
    else:
        return None
    return number if math.isfinite(number) else None


def format_ratio(number):
    text_value = ('%.4f' % number).rstrip('0').rstrip('.')
    return '0' if text_value in ('', '-0') else text_value


HEADER_NAMES = {
    'risk': {'risk', 'riskamount', 'riskusd', 'riskdollars', 'riskinusd'},
    'rr': {'rr', 'riskreward', 'riskrewardratio', 'rratio', 'rmultiple'},
    'equity': {'equity', 'balance', 'accountbalance', 'currentequity', 'accountequity'},
    'notes': {'notes', 'note', 'comment', 'comments', 'description', 'remarks'},
    'exit': {
        'exittime', 'exitdatetime', 'timeexited', 'exitedtime', 'exitedat',
        'closingtime', 'closetime', 'closedtime', 'closedat',
    },
}


def parse_column_letter(text_value):
    """'B' -> 1, 'aa' -> 26. Anything else -> None."""
    text_value = (text_value or '').strip().upper()
    if not re.fullmatch(r'[A-Z]{1,2}', text_value):
        return None
    index = 0
    for letter in text_value:
        index = index * 26 + (ord(letter) - 64)
    return index - 1


def detect_columns(rows):
    """Finds the ticker / date / time / link columns by looking at the data
    itself, so it works whether or not the sheet has a header row."""
    width = max((len(row) for row in rows), default=0)
    link_count = [0] * width
    date_count = [0] * width
    time_count = [0] * width
    upper_tickers = [set() for _ in range(width)]
    any_tickers = [set() for _ in range(width)]
    for row in rows:
        for j, (value, hyperlink) in enumerate(row):
            if clean_url(value, hyperlink):
                link_count[j] += 1
            if parse_date_cell(value):
                date_count[j] += 1
            if parse_time_cell(value):
                time_count[j] += 1
            if looks_like_ticker(value):
                any_tickers[j].add(normalize_ticker(value))
                if value.strip() == value.strip().upper():
                    upper_tickers[j].add(normalize_ticker(value))

    def most(scores):
        top = max(scores, default=0)
        return scores.index(top) if top > 0 else None

    def leftmost_common(counts):
        top = max(counts, default=0)
        if top == 0:
            return None
        return next(j for j, n in enumerate(counts) if n >= top * 0.5)

    # The ticker column is the one with the most *different* ticker-like values
    # (a Short/Long column only has two). Upper-case matches win.
    ticker_sets = upper_tickers if any(upper_tickers) else any_tickers
    columns = {
        'ticker': most([len(s) for s in ticker_sets]),
        'date': leftmost_common(date_count),
        'time': leftmost_common(time_count),
        'link': most(link_count),
    }
    # A screenshot-link column is optional (it's only needed to fill screenshots).
    required = ('ticker', 'date', 'time')
    return columns, [name for name in required if columns[name] is None]


def detect_extra_columns(rows, columns, overrides):
    """Finds Risk, R:R, Equity and Notes (needed to add new trades).

    Order of trust: column letters typed on the upload page, then header names,
    then the data itself: Risk and R:R are the two columns whose product is the
    P/L column, Equity is the right-most remaining number column, Notes is the
    column with the longest text."""
    found = {'risk': None, 'rr': None, 'equity': None, 'notes': None, 'exit': None}
    used = {index for index in columns.values() if index is not None}
    for row in rows[:10]:
        for j, (value, _) in enumerate(row):
            if isinstance(value, str):
                key = re.sub(r'[^a-z0-9]', '', value.lower())
                for name, options in HEADER_NAMES.items():
                    if key in options and found[name] is None and j not in used:
                        found[name] = j
    for name, letter in overrides.items():
        index = parse_column_letter(letter)
        if index is not None:
            found[name] = index

    # rows that look like trades (used to read the numbers)
    trade_rows = []
    for number, row in enumerate(rows, start=1):
        def at(name):
            index = columns[name]
            return row[index][0] if index < len(row) else None
        if looks_like_ticker(at('ticker')) and parse_date_cell(at('date')) and parse_time_cell(at('time')):
            trade_rows.append(number)
    sample = trade_rows[:300]
    taken = used | {i for i in found.values() if i is not None}
    width = max((len(r) for r in rows), default=0)

    numeric = {}
    for j in range(min(width, MAX_SHEET_COLUMNS)):
        if j in taken:
            continue
        values = [parse_number(rows[n - 1][j][0]) if j < len(rows[n - 1]) else None for n in sample]
        if sample and sum(v is not None for v in values) >= 0.6 * len(values):
            numeric[j] = values

    if found['risk'] is None and found['rr'] is None:
        best_matches, best = 0, None
        for r, x, p in itertools.permutations(list(numeric)[:15], 3):
            if r > x:
                continue  # risk x R:R = P/L, so r and x are interchangeable here
            total = matches = 0
            for vr, vx, vp in zip(numeric[r], numeric[x], numeric[p]):
                if vr is None or vx is None or vp is None:
                    continue
                total += 1
                if abs(vr * vx - vp) <= max(0.5, 0.01 * abs(vp)):
                    matches += 1
            if total >= 3 and matches >= 0.7 * total and matches > best_matches:
                best_matches, best = matches, (r, x, p)
        if best:
            r, x, p = best
            median = lambda col: sorted(abs(v) for v in numeric[col] if v is not None)[
                len([v for v in numeric[col] if v is not None]) // 2]
            risk, ratio = (r, x) if median(r) >= median(x) else (x, r)
            found['risk'], found['rr'] = risk, ratio
            taken |= {r, x, p}
            numeric = {j: v for j, v in numeric.items() if j not in taken}

    if found['equity'] is None:
        candidates = [
            j for j, values in numeric.items()
            if j not in taken and sum(v is not None for v in values) >= 0.8 * len(values)
        ]
        if candidates:
            found['equity'] = max(candidates)
    taken |= {i for i in found.values() if i is not None}

    if found['notes'] is None:
        best_length, best_col = 0, None
        for j in range(min(width, MAX_SHEET_COLUMNS)):
            if j in taken:
                continue
            lengths = [
                len(rows[n - 1][j][0].strip())
                for n in sample
                if j < len(rows[n - 1]) and isinstance(rows[n - 1][j][0], str)
                and not clean_url(rows[n - 1][j][0], rows[n - 1][j][1])
            ]
            if lengths and sum(lengths) / len(lengths) > best_length:
                best_length, best_col = sum(lengths) / len(lengths), j
        if best_length >= 12:
            found['notes'] = best_col
    return found


def parse_sheet_rows(rows, columns):
    """columns holds every column index (ticker, date, time, link, risk, rr,
    equity, notes); the last four may be None."""
    entries = []
    for number, row in enumerate(rows, start=1):
        def cell(name):
            index = columns.get(name)
            if index is None or index >= len(row):
                return (None, None)
            return row[index]

        ticker = cell('ticker')[0]
        date_parts = parse_date_cell(cell('date')[0])
        time_parts = parse_time_cell(cell('time')[0])
        if not looks_like_ticker(ticker) or not date_parts or not time_parts:
            continue  # header, blank or unrelated row
        link_value, link_target = cell('link')
        notes = cell('notes')[0]
        entries.append({
            'row': number,
            'ticker': normalize_ticker(ticker),
            'year': date_parts[0],
            'month': date_parts[1],
            'day': date_parts[2],
            'hour': time_parts[0],
            'minute': time_parts[1],
            'url': clean_url(link_value, link_target),
            'risk': parse_number(cell('risk')[0]),
            'rr': parse_number(cell('rr')[0]),
            'equity': parse_number(cell('equity')[0]),
            'notes': notes.strip() if isinstance(notes, str) else '',
            'exit_text': (
                format_exit_datetime(exit_when)
                if (exit_when := parse_exit_cell(cell('exit')[0])) else None
            ),
        })
    return entries


def describe_excel_time(entry):
    text_value = datetime(2000, entry['month'], entry['day']).strftime('%d %b')
    if entry['year']:
        text_value += f" {entry['year']}"
    return f"{text_value} {entry['hour']:02d}:{entry['minute']:02d}"


def entry_datetime_text(entry):
    """The app's own format, e.g. 2026-10-06T11:16. A sheet date without a
    year gets this year (or last year if that would be in the future)."""
    year = entry['year']
    if year is None:
        today = datetime.now()
        year = today.year
        try:
            if datetime(year, entry['month'], entry['day']) > today + timedelta(days=7):
                year -= 1
        except ValueError:
            year -= 1  # 29 Feb in a non-leap year
    return f"{year:04d}-{entry['month']:02d}-{entry['day']:02d}T{entry['hour']:02d}:{entry['minute']:02d}"


def setup_from_notes(notes):
    """Best guess of Good / Medium / Bad from the trade's note. 'bad' wins, then
    'good'; 'decent' or no clue at all means Medium. It's only a starting point:
    the preview lets you change it for every trade."""
    note = (notes or '').lower()
    if re.search(r'\bbad\b|not a good|not good|forcing|\bpoor\b', note):
        return 'bad'
    if re.search(r'\b(perfect|good|great|excellent)\b', note):
        return 'good'
    return 'medium'


def build_import_preview(account, sheets, overrides):
    """Works out what the sheet(s) can fill in and which rows are new."""
    trades = (
        Trade.query.filter_by(account_id=account.id)
        .order_by(Trade.position.asc(), Trade.id.asc())
        .all()
    )
    index = {}
    for trade in trades:
        when = parse_entry_datetime(trade.entry_datetime)
        if when:
            key = (normalize_ticker(trade.ticker), when.month, when.day, when.hour, when.minute)
            index.setdefault(key, []).append((trade, when))

    sheet_notes, unmatched, kept_different, proposals, new_trades = [], [], [], {}, []
    exit_proposals, exit_already = {}, 0
    seen_new = set()
    rows_read = rows_without_link = already_same = 0
    for name, rows in sheets:
        columns, missing = detect_columns(rows)
        if missing:
            sheet_notes.append({
                'name': name, 'columns': None, 'extra': None,
                'problem': 'Skipped: could not find '
                + ', '.join(COLUMN_NAMES[m] for m in missing) + '.',
            })
            continue
        extra = detect_extra_columns(rows, columns, overrides)
        entries = parse_sheet_rows(rows, {**columns, **extra})
        letters = lambda cols: {k: (column_letter(v) if v is not None else None) for k, v in cols.items()}
        sheet_notes.append({
            'name': name, 'columns': letters(columns), 'extra': letters(extra),
            'rows': len(entries), 'problem': None,
            'cant_add': [label for label, key in (('Risk', 'risk'), ('R:R', 'rr'), ('Equity', 'equity')) if extra[key] is None],
        })
        rows_read += len(entries)
        for entry in entries:
            entry['sheet'] = name
            key = (entry['ticker'], entry['month'], entry['day'], entry['hour'], entry['minute'])
            candidates = [
                trade for trade, when in index.get(key, [])
                if entry['year'] is None or when.year == entry['year']
            ]
            shown = {
                'sheet': name, 'row': entry['row'], 'ticker': entry['ticker'],
                'when': describe_excel_time(entry), 'url': entry['url'],
            }
            if len(candidates) > 1:
                unmatched.append(dict(shown, reason=f'{len(candidates)} trades match this row, so it is ambiguous.'))
            elif len(candidates) == 1:
                trade = candidates[0]
                if entry['exit_text']:
                    if (trade.exit_datetime or '').strip():
                        exit_already += 1
                    else:
                        exit_proposals.setdefault(trade.id, {
                            'trade': trade, 'exit': entry['exit_text'],
                            'row': entry['row'], 'sheet': name,
                        })
                if not entry['url']:
                    rows_without_link += 1
                elif has_screenshot(trade):
                    if trade.screenshot_url.strip() == entry['url']:
                        already_same += 1
                    else:
                        kept_different.append({'trade': trade, 'url': entry['url']})
                else:
                    proposals.setdefault(trade.id, {'trade': trade, 'entries': []})['entries'].append(entry)
            else:  # not in the app
                lacking = [label for label, key2 in (('Risk', 'risk'), ('R:R', 'rr'), ('Equity', 'equity')) if entry[key2] is None or (key2 == 'risk' and entry[key2] < 0)]
                if lacking:
                    unmatched.append(dict(shown, reason="Not in the app, and can't be added: no usable "
                                          + ', '.join(lacking) + ' in this row.'))
                elif key in seen_new:
                    unmatched.append(dict(shown, reason='Appears twice in the sheet; only the first one is offered.'))
                else:
                    seen_new.add(key)
                    when_text = entry_datetime_text(entry)
                    quality = setup_from_notes(entry['notes'])
                    new_trades.append({
                        'ticker': entry['ticker'], 'when': when_text,
                        'when_text': describe_excel_time(dict(entry, year=int(when_text[:4]))),
                        'risk': entry['risk'], 'rr': entry['rr'], 'equity': entry['equity'],
                        'pl': entry['risk'] * entry['rr'],
                        'url': entry['url'], 'notes': entry['notes'], 'quality': quality,
                        'exit_text': (entry['exit_text'] or '').replace('T', ' '),
                        'sheet': name, 'row': entry['row'],
                        'payload': json.dumps({
                            'ticker': entry['ticker'], 'when': when_text, 'risk': entry['risk'],
                            'rr': entry['rr'], 'equity': entry['equity'], 'url': entry['url'] or '',
                            'exit': entry['exit_text'] or '',
                        }, separators=(',', ':')),
                    })

    fills = []
    for item in proposals.values():
        entries = item['entries']
        if len({e['url'] for e in entries}) > 1:
            first = entries[0]
            unmatched.append({
                'sheet': first['sheet'], 'row': first['row'], 'ticker': first['ticker'],
                'when': describe_excel_time(first), 'url': first['url'],
                'reason': 'The sheet has different links for this one trade (rows '
                + ', '.join(str(e['row']) for e in entries) + '), so it was left out.',
            })
        else:
            fills.append({
                'trade': item['trade'], 'url': entries[0]['url'],
                'row': entries[0]['row'], 'sheet': entries[0]['sheet'],
            })
    fills.sort(key=lambda f: (f['trade'].position or 0, f['trade'].id))
    exit_fills = sorted(
        exit_proposals.values(), key=lambda f: (f['trade'].position or 0, f['trade'].id)
    )
    new_trades.sort(key=lambda n: n['when'])
    filled_ids = {f['trade'].id for f in fills}
    still_missing = [t for t in trades if not has_screenshot(t) and t.id not in filled_ids]
    return {
        'sheets': sheet_notes,
        'rows_read': rows_read,
        'rows_without_link': rows_without_link,
        'fills': fills,
        'exit_fills': exit_fills,
        'exit_already': exit_already,
        'new_trades': new_trades[:MAX_NEW_TRADES],
        'new_total': len(new_trades),
        'unmatched': unmatched[:MAX_LISTED_ROWS],
        'unmatched_total': len(unmatched),
        'kept_different': kept_different,
        'already_same': already_same,
        'still_missing': still_missing,
    }


def answer_points():
    """Points 'Yes' / 'No' are worth for the two questions a sheet can't answer."""
    return {
        'risk': get_rule_points('risk_under_half'),
        'daily': get_rule_points('daily_dd'),
    }


@app.route('/import-screenshots', methods=['GET', 'POST'])
def import_screenshots():
    account = get_active_account()
    trades = Trade.query.filter_by(account_id=account.id).all()
    context = {
        'stage': 'upload',
        'account': account,
        'total_trades': len(trades),
        'missing_now': sum(1 for t in trades if not has_screenshot(t)),
        'error': None,
        'result': None,
        'points': answer_points(),
        'filled': request.args.get('filled', type=int),
        'skipped': request.args.get('skipped', type=int),
        'added': request.args.get('added', type=int),
        'add_skipped': request.args.get('add_skipped', type=int),
        'exits': request.args.get('exits', type=int),
        'exit_skipped': request.args.get('exit_skipped', type=int),
    }
    if request.method == 'POST':
        upload = request.files.get('file')
        overrides = {
            'risk': request.form.get('risk_col', ''),
            'rr': request.form.get('rr_col', ''),
            'equity': request.form.get('equity_col', ''),
        }
        try:
            for label, letter in overrides.items():
                if letter.strip() and parse_column_letter(letter) is None:
                    raise UploadError(f'"{letter.strip()}" is not a column letter. Use letters like B or AA.')
            if upload is None or not upload.filename:
                raise UploadError('Please choose an Excel (.xlsx) or CSV file first.')
            data = upload.read(MAX_UPLOAD_BYTES + 1)
            if len(data) > MAX_UPLOAD_BYTES:
                raise UploadError('That file is too large (the limit is 10 MB).')
            context['result'] = build_import_preview(
                account, read_sheets(upload.filename, data), overrides
            )
            context['stage'] = 'preview'
        except UploadError as problem:
            context['error'] = str(problem)
    return render_template('import_screenshots.html', **context)


def fill_screenshots(account):
    """Step 1 of confirming: fill the ticked, still-empty screenshots."""
    filled = skipped = 0
    for raw_id in dict.fromkeys(request.form.getlist('apply')):
        try:
            trade = db.session.get(Trade, int(raw_id))
        except ValueError:
            trade = None
        url = clean_url(request.form.get(f'url_{raw_id}'))
        # Re-check everything: only this account, only still-empty screenshots,
        # only real http(s) links. Existing screenshots are never overwritten.
        if trade is None or trade.account_id != account.id or has_screenshot(trade) or not url:
            skipped += 1
            continue
        trade.screenshot_url = url
        filled += 1
    return filled, skipped


def fill_exit_times(account):
    """Fill the ticked exit times. Only this account's trades that have no exit
    time yet are touched: a saved exit time is never overwritten."""
    filled = skipped = 0
    for raw_id in dict.fromkeys(request.form.getlist('apply_exit')):
        try:
            trade = db.session.get(Trade, int(raw_id))
        except ValueError:
            trade = None
        exit_text = clean_exit_datetime(request.form.get(f'exit_{raw_id}'))
        if (
            trade is None or trade.account_id != account.id
            or (trade.exit_datetime or '').strip() or not exit_text
        ):
            skipped += 1
            continue
        trade.exit_datetime = exit_text
        filled += 1
    return filled, skipped


def add_new_trades(account):
    """Step 2 of confirming: create the ticked rows as new trades."""
    answers = {}
    for field in ('answer_risk', 'answer_daily'):
        choice = request.form.get(field, 'no')
        answers[field] = choice if choice in ('yes', 'no', 'blank') else 'no'

    ordered = [
        (t, parse_entry_datetime(t.entry_datetime))
        for t in Trade.query.filter_by(account_id=account.id)
        .order_by(Trade.position.asc(), Trade.id.asc()).all()
    ]
    keys = {
        (normalize_ticker(t.ticker), w.month, w.day, w.hour, w.minute)
        for t, w in ordered if w
    }
    added = skipped = 0
    additions = []
    for raw_index in list(dict.fromkeys(request.form.getlist('add')))[:MAX_NEW_TRADES]:
        try:
            row = json.loads(request.form.get(f'row_{raw_index}', ''))
            ticker = normalize_ticker(row['ticker'])
            when = parse_entry_datetime(row['when'])
            risk, ratio, equity = (parse_number(row[k]) for k in ('risk', 'rr', 'equity'))
        except (ValueError, KeyError, TypeError):
            skipped += 1
            continue
        key = (ticker, when.month, when.day, when.hour, when.minute) if when else None
        if (
            not TICKER_PATTERN.match(ticker) or when is None
            or risk is None or risk < 0 or ratio is None or equity is None
            or abs(risk) > 1e12 or abs(ratio) > 1e6 or abs(equity) > 1e12
            or key in keys  # already in the app (e.g. this form was sent twice)
        ):
            skipped += 1
            continue
        quality = request.form.get(f'quality_{raw_index}')
        if quality not in ('good', 'medium', 'bad'):
            quality = 'medium'
        risk_answer = None if answers['answer_risk'] == 'blank' else answers['answer_risk']
        daily_answer = None if answers['answer_daily'] == 'blank' else answers['answer_daily']
        scores = calculate_scores({
            'setup_quality': quality, 'profit_taken': 'no', 'stop_loss_zero': 'no',
            'risk_under_half': risk_answer, 'daily_risk': daily_answer,
            'risk_reward_val': ratio, 'overtrading': 'no', 'revenge_trading': 'no',
            'impatient_trading': 'no', 'fearful_trading': 'no',
        })
        trade = Trade(
            account_id=account.id, position=0, setup_quality_str=quality, ticker=ticker,
            entry_datetime=f'{when:%Y-%m-%dT%H:%M}', profit_taken_str='no', stop_loss_zero_str='no',
            current_equity=equity, risk_str=format_number(risk), risk_reward_str=format_ratio(ratio),
            risk_under_half_str=risk_answer, daily_risk_str=daily_answer,
            overtrading_str='no', revenge_trading_str='no',
            impatient_trading_str='no', fearful_trading_str='no',
            screenshot_url=clean_url(row.get('url')),
            exit_datetime=clean_exit_datetime(row.get('exit')),
            setup_score=scores['setup_score'], profit_score=scores['profit_score'],
            stop_loss_score=scores['stop_loss_score'], risk_score=scores['risk_score'],
            daily_risk_score=scores['daily_risk_score'], psychology_score=scores['psychology_score'],
            risk_reward_score=scores['risk_reward_score'], total_score=scores['total_score'],
        )
        db.session.add(trade)
        keys.add(key)
        additions.append((when, trade))
        added += 1

    # Put each new trade where it belongs in time, so the charts stay in order.
    for when, trade in sorted(additions, key=lambda a: a[0]):
        position = 0
        for i, (other, other_when) in enumerate(ordered):
            if other_when is not None and other_when <= when:
                position = i + 1
        ordered.insert(position, (trade, when))
    if additions:
        for number, (trade, _) in enumerate(ordered, start=1):
            trade.position = number
    return added, skipped


@app.route('/import-screenshots/apply', methods=['POST'])
def apply_import():
    account = get_active_account()
    filled, skipped = fill_screenshots(account)
    exits, exit_skipped = fill_exit_times(account)
    added, add_skipped = add_new_trades(account)
    db.session.commit()
    return redirect(url_for(
        'import_screenshots', filled=filled, skipped=skipped,
        added=added, add_skipped=add_skipped,
        exits=exits, exit_skipped=exit_skipped,
    ))


# --- TRADING QUALITY APP ROUTES ---


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/new-trade')
def new_trade_page():
    pts = {key: get_rule_points(key) for key in SCORING_RULES}
    account = get_active_account()
    has_trades = Trade.query.filter_by(account_id=account.id).first() is not None
    # First trade in a brand-new account: start from that account's own balance.
    default_equity = (
        account.starting_balance
        if not has_trades and account.starting_balance is not None
        else None
    )
    return render_template('new_trade.html', pts=pts, default_equity=default_equity)


@app.route('/trades-list')
def trades_list():
    sort_mode = 'position' if request.args.get('sort') == 'position' else 'exit'
    return render_template(
        'trades.html', trades=ordered_trades(sort_mode), sort_mode=sort_mode
    )


def ordered_trades(sort_mode):
    """The active account's trades. Default ('exit'): most recently closed
    first, a trade with no exit time placed by its entry time. 'position' is the
    manual order (the one the move up/down arrows change)."""
    all_trades = (
        Trade.query.filter_by(account_id=get_active_account().id)
        .order_by(Trade.position.asc())
        .all()
    )
    if sort_mode == 'exit':
        all_trades.sort(
            key=lambda t: (
                (t.exit_datetime or t.entry_datetime or '').replace(' ', 'T'),
                t.position or 0,
            ),
            reverse=True,
        )
    return all_trades


def excel_number(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


@app.route('/trades-list/export')
def export_trades():
    """Download the trades list (same order as on screen) as an .xlsx file."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    account = get_active_account()
    sort_mode = 'position' if request.args.get('sort') == 'position' else 'exit'
    trades = ordered_trades(sort_mode)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Trades'
    headers = [
        'ID', 'Ticker', 'Entry Date/Time', 'Exit Date/Time', 'Setup', '1/3 Profit',
        'SL Zero', 'Risk ($)', 'R:R', 'Profit / Loss ($)', 'Risk < 0.5%',
        'Daily DD < 1.5%', 'Psychology', 'Equity', 'Total Score', 'Chart',
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    def put_text(row, col, value):
        cell = ws.cell(row=row, column=col, value=value)
        if isinstance(value, str) and value[:1] in ('=', '+', '-', '@'):
            cell.data_type = 's'  # never let a ticker or link become a formula
        return cell

    for r, t in enumerate(trades, start=2):
        risk = excel_number(t.risk_str)
        rr = excel_number(t.risk_reward_str)
        entry = parse_entry_datetime(t.entry_datetime)
        exit_when = parse_entry_datetime(t.exit_datetime)
        put_text(r, 1, t.id)
        put_text(r, 2, t.ticker)
        for col, when in ((3, entry), (4, exit_when)):
            cell = ws.cell(row=r, column=col, value=when)
            cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        put_text(r, 5, t.setup_quality_str)
        put_text(r, 6, (t.profit_taken_str or '').capitalize())
        put_text(r, 7, (t.stop_loss_zero_str or '').capitalize())
        ws.cell(row=r, column=8, value=risk)
        ws.cell(row=r, column=9, value=rr)
        ws.cell(row=r, column=10, value=(risk or 0) * (rr or 0)).number_format = '0.00'
        put_text(r, 11, (t.risk_under_half_str or '').capitalize())
        put_text(r, 12, (t.daily_risk_str or '').capitalize())
        put_text(r, 13, ', '.join(t.psychology_flags))
        ws.cell(row=r, column=14, value=t.current_equity).number_format = '0.00'
        ws.cell(row=r, column=15, value=t.total_score)
        if t.screenshot_url:
            cell = put_text(r, 16, t.screenshot_url)
            if t.screenshot_url.lower().startswith(('http://', 'https://')):
                cell.hyperlink = t.screenshot_url

    widths = [7, 12, 20, 20, 10, 10, 9, 10, 8, 16, 12, 15, 22, 12, 12, 45]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(len(trades) + 1, 2)}'

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    safe_name = re.sub(r'[^A-Za-z0-9_-]+', '_', account.name or 'trades').strip('_') or 'trades'
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f'{safe_name}_trades_{date_type.today():%Y-%m-%d}.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/trade/edit/<int:id>', methods=['GET', 'POST'])
def edit_trade(id):
    trade = Trade.query.get_or_404(id)
    if request.method == 'POST':
        data = request.form
        scores = calculate_scores({
            'setup_quality': data.get('setup_quality'),
            'profit_taken': data.get('profit_taken'),
            'stop_loss_zero': data.get('stop_loss_zero'),
            'risk_reward_val': data.get('risk_reward_val'),
            'risk_under_half': data.get('risk_under_half'),
            'daily_risk': data.get('daily_risk'),
            'overtrading': data.get('overtrading'),
            'revenge_trading': data.get('revenge_trading'),
            'impatient_trading': data.get('impatient_trading'),
            'fearful_trading': data.get('fearful_trading'),
        })

        trade.setup_quality_str = data.get('setup_quality')
        trade.ticker = data.get('ticker')
        trade.entry_datetime = data.get('entry_datetime')
        trade.exit_datetime = clean_exit_datetime(data.get('exit_datetime'))
        trade.profit_taken_str = data.get('profit_taken')
        trade.stop_loss_zero_str = data.get('stop_loss_zero')
        trade.current_equity = float(data.get('current_equity', 0))
        trade.risk_str = str(data.get('risk_val'))
        trade.risk_reward_str = str(data.get('risk_reward_val'))

        apply_scores(trade, scores)
        trade.risk_under_half_str = data.get('risk_under_half')
        trade.daily_risk_str = data.get('daily_risk')
        trade.overtrading_str = data.get('overtrading', 'no')
        trade.revenge_trading_str = data.get('revenge_trading', 'no')
        trade.impatient_trading_str = data.get('impatient_trading', 'no')
        trade.fearful_trading_str = data.get('fearful_trading', 'no')

        trade.screenshot_url = request.form.get('screenshot_url', '').strip() or None

        db.session.commit()
        return redirect(url_for('trades_list'))

    return render_template('edit_trade.html', trade=trade)


@app.route('/trade/delete/<int:id>', methods=['POST', 'GET'])
def delete_trade(id):
    trade = Trade.query.get_or_404(id)
    db.session.delete(trade)
    db.session.commit()
    return redirect(url_for('trades_list'))


@app.route('/api/trades', methods=['GET', 'POST'])
def manage_trades():
    if request.method == 'POST':
        data = request.json
        scores = calculate_scores(data)
        account = get_active_account()
        max_pos = (
            db.session.query(db.func.max(Trade.position))
            .filter(Trade.account_id == account.id)
            .scalar()
            or 0
        )
        new_trade = Trade(
            account_id=account.id,
            position=max_pos + 1,
            setup_quality_str=data.get('setup_quality'),
            ticker=data.get('ticker'),
            entry_datetime=data.get('entry_datetime'),
            exit_datetime=clean_exit_datetime(data.get('exit_datetime')),
            profit_taken_str=data.get('profit_taken'),
            stop_loss_zero_str=data.get('stop_loss_zero'),
            current_equity=float(data.get('current_equity', 0)),
            risk_str=str(data.get('risk_val')),
            risk_reward_str=str(data.get('risk_reward_val')),
            risk_under_half_str=data.get('risk_under_half'),
            daily_risk_str=data.get('daily_risk'),
            overtrading_str=data.get('overtrading', 'no'),
            revenge_trading_str=data.get('revenge_trading', 'no'),
            impatient_trading_str=data.get('impatient_trading', 'no'),
            fearful_trading_str=data.get('fearful_trading', 'no'),
            screenshot_url=data.get('screenshot_url', '').strip() or None,
            setup_score=scores['setup_score'],
            profit_score=scores['profit_score'],
            stop_loss_score=scores['stop_loss_score'],
            risk_score=scores['risk_score'],
            daily_risk_score=scores['daily_risk_score'],
            psychology_score=scores['psychology_score'],
            risk_reward_score=scores['risk_reward_score'],
            total_score=scores['total_score'],
        )
        db.session.add(new_trade)
        db.session.commit()
        return jsonify({
            'status': 'success',
            'total_score': scores['total_score'],
        })

    trades = (
        Trade.query.filter_by(account_id=get_active_account().id)
        .order_by(Trade.position.asc())
        .all()
    )
    result = []
    for t in trades:
        result.append({
            'id': t.id,
            'ticker': t.ticker,
            'entry_datetime': t.entry_datetime,
            'exit_datetime': t.exit_datetime,
            'current_equity': t.current_equity,
            'risk_val': float(t.risk_str) if t.risk_str else 0.0,
            'risk_reward_str': (
                float(t.risk_reward_str) if t.risk_reward_str else 0.0
            ),
            'total_score': t.total_score,
        })
    return jsonify(result)


@app.route('/trade/move/<int:id>/<direction>')
def move_trade(id, direction):
    trade = Trade.query.get_or_404(id)

    if direction == 'up':
        adjacent = (
            Trade.query.filter(
                Trade.account_id == trade.account_id,
                Trade.position < trade.position,
            )
            .order_by(Trade.position.desc())
            .first()
        )
    elif direction == 'down':
        adjacent = (
            Trade.query.filter(
                Trade.account_id == trade.account_id,
                Trade.position > trade.position,
            )
            .order_by(Trade.position.asc())
            .first()
        )
    else:
        adjacent = None

    if adjacent:
        trade.position, adjacent.position = adjacent.position, trade.position
        db.session.commit()

    return redirect(url_for('trades_list', sort='position'))


# --- IBKR TRADES ROUTES ---


@app.route('/ibkr', methods=['GET', 'POST'])
def ibkr_trades():
    if request.method == 'POST':
        try:
            entry_price_val = float(request.form.get('entry_price', 0))
        except ValueError:
            entry_price_val = 0.0

        try:
            exit_price_val = float(request.form.get('exit_price', 0))
        except ValueError:
            exit_price_val = 0.0

        try:
            pnl_val = float(request.form.get('pnl', 0))
        except ValueError:
            pnl_val = 0.0

        try:
            balance_val = float(request.form.get('balance', 0))
        except ValueError:
            balance_val = 0.0

        new_ibkr_trade = IBKRTrade(
            ticker=request.form.get('ticker'),
            time_entered=request.form.get('time_entered'),
            time_exited=request.form.get('time_exited'),
            entry_price=entry_price_val,
            exit_price=exit_price_val,
            pnl=pnl_val,
            balance=balance_val,
        )
        db.session.add(new_ibkr_trade)
        db.session.commit()
        return redirect(url_for('ibkr_trades'))

    all_ibkr_trades = IBKRTrade.query.order_by(IBKRTrade.id.asc()).all()
    return render_template('ibkr.html', trades=all_ibkr_trades)


@app.route('/ibkr/delete/<int:id>', methods=['POST', 'GET'])
def delete_ibkr_trade(id):
    ibkr_trade = IBKRTrade.query.get_or_404(id)
    db.session.delete(ibkr_trade)
    db.session.commit()
    return redirect(url_for('ibkr_trades'))


@app.route('/ibkr/edit/<int:id>', methods=['GET', 'POST'])
def edit_ibkr_trade(id):
    ibkr_trade = IBKRTrade.query.get_or_404(id)
    if request.method == 'POST':
        try:
            ibkr_trade.ticker = request.form.get('ticker')
            ibkr_trade.time_entered = request.form.get('time_entered')
            ibkr_trade.time_exited = request.form.get('time_exited')
            ibkr_trade.entry_price = float(request.form.get('entry_price', 0))
            ibkr_trade.exit_price = float(request.form.get('exit_price', 0))
            ibkr_trade.pnl = float(request.form.get('pnl', 0))
            ibkr_trade.balance = float(request.form.get('balance', 0))
            db.session.commit()
            return redirect(url_for('ibkr_trades'))
        except ValueError:
            pass

    return render_template('edit_ibkr.html', trade=ibkr_trade)


@app.route('/ibkr-charts')
def ibkr_charts():
    return render_template('ibkr_charts.html')


@app.route('/api/ibkr-trades')
def api_ibkr_trades():
    trades = IBKRTrade.query.order_by(IBKRTrade.id.asc()).all()
    return jsonify([
        {
            'id': t.id,
            'ticker': t.ticker,
            'time_entered': t.time_entered,
            'time_exited': t.time_exited,
            'entry_price': t.entry_price,
            'exit_price': t.exit_price,
            'pnl': t.pnl,
            'balance': t.balance,
        }
        for t in trades
    ])


# --- SETTINGS ROUTES ---

SETTINGS_PAGE_RULES = {
    1: ('Setup Quality', 'Good / Medium / Bad', 'setup_good'),
    2: ('1/3rd Profit Taken', 'Yes / No', 'profit_yes'),
    3: ('Stop Loss Moved to Zero', 'Yes / No', 'sl_yes'),
    4: ('Risk < 0.5%', 'Yes / No', 'risk_under_half'),
    5: ('Daily Drawdown < 1.5%', 'Yes / No', 'daily_dd'),
    6: (
        'Risk:Reward Ratio (R:R)',
        '<0 / 0-0.5 / 0.5-1 / 1-2 / 2-3 / 3+',
        'rr_bands',
    ),
    7: ('Overtrading', 'Yes / No', 'overtrading_yes'),
    8: ('Revenge Trading', 'Yes / No', 'revenge_yes'),
    9: ('Impatient Trading', 'Yes / No', 'impatient_yes'),
    10: ('Fearful Trading', 'Yes / No', 'fearful_yes'),
}


def current_points_text(key):
    return format_points(get_rule_points(key))


@app.route('/settings')
def settings_page():
    scoring_rules = [
        {
            'id': rule_id,
            'category': category,
            'condition': condition,
            'key': key,
            'points': current_points_text(key),
        }
        for rule_id, (category, condition, key) in SETTINGS_PAGE_RULES.items()
    ]
    return render_template('settings.html', rules=scoring_rules)


@app.route('/settings/edit/<int:id>', methods=['GET', 'POST'])
def edit_setting(id):
    if id not in SETTINGS_PAGE_RULES:
        return redirect(url_for('settings_page'))

    category, _condition, key = SETTINGS_PAGE_RULES[id]
    rule_info = SCORING_RULES[key]
    rule = {'category': category, 'points': current_points_text(key)}
    error = None

    if request.method == 'POST':
        submitted = (request.form.get('points') or '').strip()
        parsed = parse_points(submitted, rule_info['count'])

        if parsed is None:
            error = (
                f"Please enter exactly {rule_info['count']} numbers separated by"
                f" '/' or ','. {rule_info['hint']}."
            )
            rule['points'] = submitted
        else:
            value = format_points(parsed)
            setting_obj = AppSetting.query.filter_by(key=key).first()
            if setting_obj:
                setting_obj.value = value
            else:
                db.session.add(AppSetting(key=key, value=value))
            db.session.commit()

            recalculate_all_trades()
            return redirect(url_for('settings_page'))

    return render_template(
        'edit_setting.html', rule=rule, id=id, error=error, hint=rule_info['hint']
    )


if __name__ == '__main__':
    app.run(debug=True)